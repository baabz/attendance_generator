"""
Generator for Official Kaduna Polytechnic Exam Attendance Sheets in Microsoft Word (.docx) format.
"""
import io
from typing import List, Union
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from .models import ExamConfig, VenueAllocation


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Set inner padding for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def make_row_header(row):
    """Mark a table row to repeat as header across page splits."""
    trPr = row._tr.get_or_add_trPr()
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    trPr.append(tblHeader)


def prevent_row_split(row):
    """Prevent row from breaking across pages."""
    trPr = row._tr.get_or_add_trPr()
    cantSplit = parse_xml(f'<w:cantSplit {nsdecls("w")}/>')
    trPr.append(cantSplit)


def build_attendance_docx(
    config: ExamConfig,
    allocations: List[VenueAllocation],
    output_target: Union[str, io.BytesIO, None] = None
) -> bytes:
    """
    Build a single Word (.docx) file containing all allocated halls.
    Each hall begins on a new page with the full institution header and student roster.
    """
    doc = Document()

    # Configure Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base font setting
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    col_widths = [
        Inches(0.55),  # S/N
        Inches(1.45),  # REG NO
        Inches(2.55),  # NAMES
        Inches(1.05),  # BOOKLET NO
        Inches(0.80),  # SIGN IN
        Inches(0.80),  # SIGN OUT
    ]

    running_sn = 1

    for hall_idx, alloc in enumerate(allocations):
        # Insert page break between halls
        if hall_idx > 0:
            doc.add_page_break()

        # Institution Header
        p_sem = doc.add_paragraph()
        p_sem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sem.paragraph_format.space_before = Pt(0)
        p_sem.paragraph_format.space_after = Pt(2)
        r = p_sem.add_run(config.session_semester.strip())
        r.bold = True
        r.font.size = Pt(11)

        p_dept = doc.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dept.paragraph_format.space_before = Pt(0)
        p_dept.paragraph_format.space_after = Pt(2)
        r = p_dept.add_run(config.department.strip())
        r.bold = True
        r.font.size = Pt(11)

        p_col = doc.add_paragraph()
        p_col.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_col.paragraph_format.space_before = Pt(0)
        p_col.paragraph_format.space_after = Pt(2)
        r = p_col.add_run(config.college.strip())
        r.bold = True
        r.font.size = Pt(11)

        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(2)
        p_title.paragraph_format.space_after = Pt(6)
        r = p_title.add_run("ATTENDANCE SHEET")
        r.bold = True
        r.font.size = Pt(12)

        # Metadata grid using a borderless 2-column layout for pixel-perfect alignment
        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        left_width = Inches(4.3)
        right_width = Inches(2.9)
        for row in meta_table.rows:
            row.cells[0].width = left_width
            row.cells[1].width = right_width

        meta_rows_data = [
            (f"COURSE: {config.course_title.strip()}", f"CODE: {config.course_code.strip()}"),
            (f"HALL: {alloc.venue_name.strip()}", f"DATE: {config.exam_date.strip()}"),
            (
                f"TIME IN: {config.time_in.strip() if config.time_in.strip() else '...................................'}",
                f"TIME OUT: {config.time_out.strip() if config.time_out.strip() else '...................................'}"
            ),
            (
                f"LEVEL: {config.level.strip()}",
                f"Chief Invigilator: {config.chief_invigilator.strip() if config.chief_invigilator.strip() else '____________________'}"
            ),
        ]

        for i, (left_text, right_text) in enumerate(meta_rows_data):
            c_left = meta_table.rows[i].cells[0]
            c_right = meta_table.rows[i].cells[1]

            p_l = c_left.paragraphs[0]
            p_l.paragraph_format.space_before = Pt(1)
            p_l.paragraph_format.space_after = Pt(1)
            r_l = p_l.add_run(left_text)
            r_l.bold = True
            r_l.font.size = Pt(10)

            p_r = c_right.paragraphs[0]
            p_r.paragraph_format.space_before = Pt(1)
            p_r.paragraph_format.space_after = Pt(1)
            r_r = p_r.add_run(right_text)
            r_r.bold = True
            r_r.font.size = Pt(10)

        # Invigilators lines (Right aligned)
        invig_count = config.invigilator_count or 3
        for num in range(1, invig_count + 1):
            p_inv = doc.add_paragraph()
            p_inv.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_inv.paragraph_format.space_before = Pt(0)
            p_inv.paragraph_format.space_after = Pt(1)
            r_num = p_inv.add_run(f"{num}. .....................................................................")
            r_num.font.size = Pt(9.5)

        # Space before table
        p_gap = doc.add_paragraph()
        p_gap.paragraph_format.space_before = Pt(0)
        p_gap.paragraph_format.space_after = Pt(4)

        # Student Table
        student_rows_count = len(alloc.students) + 1  # +1 for header
        table = doc.add_table(rows=student_rows_count, cols=6)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ['S/N', 'REG NO', 'NAMES', 'BOOKLET NO', 'SIGN IN', 'SIGN OUT']

        # Format Header Row
        hdr_row = table.rows[0]
        make_row_header(hdr_row)
        prevent_row_split(hdr_row)

        for col_idx, text in enumerate(headers):
            cell = hdr_row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(9.5)

        # Format Student Rows
        for s_idx, student in enumerate(alloc.students):
            row_idx = s_idx + 1
            row = table.rows[row_idx]
            prevent_row_split(row)

            sn_val = str(running_sn) if config.continuous_sn else str(s_idx + 1)
            if config.continuous_sn:
                running_sn += 1

            cell_values = [
                sn_val,
                student.reg_no,
                student.name,
                '',  # Booklet No (blank)
                '',  # Sign In (blank)
                ''   # Sign Out (blank)
            ]

            for col_idx, val in enumerate(cell_values):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3, 4, 5] else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.size = Pt(9)

    # Save to buffer or file
    if output_target is None:
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    elif isinstance(output_target, str):
        doc.save(output_target)
        with open(output_target, 'rb') as f:
            return f.read()
    elif hasattr(output_target, 'write'):
        doc.save(output_target)
        if hasattr(output_target, 'getvalue'):
            return output_target.getvalue()
        return b''
