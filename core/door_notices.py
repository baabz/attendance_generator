"""
Door Notices / Venue Entrance Slips Generator.
Produces high-visibility, print-ready door notices for examination halls in Word (.docx) and PDF (.pdf).
"""
import io
from typing import List, Union
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from .models import ExamConfig, VenueAllocation


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Set inner padding for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_door_notices_docx(
    config: ExamConfig,
    allocations: List[VenueAllocation],
    output_target: Union[str, io.BytesIO, None] = None
) -> bytes:
    """
    Generate official Hall Door Notices in Word format.
    Each hall has a prominent headline banner and candidate roster.
    """
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    for hall_idx, alloc in enumerate(allocations):
        if hall_idx > 0:
            doc.add_page_break()

        # Institutional Header
        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_inst.paragraph_format.space_before = Pt(0)
        p_inst.paragraph_format.space_after = Pt(1)
        r = p_inst.add_run(config.college.strip().upper())
        r.bold = True
        r.font.size = Pt(11)

        p_dept = doc.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dept.paragraph_format.space_before = Pt(0)
        p_dept.paragraph_format.space_after = Pt(1)
        r = p_dept.add_run(config.department.strip().upper())
        r.bold = True
        r.font.size = Pt(10.5)

        p_sem = doc.add_paragraph()
        p_sem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sem.paragraph_format.space_before = Pt(0)
        p_sem.paragraph_format.space_after = Pt(4)
        r = p_sem.add_run(config.session_semester.strip())
        r.bold = True
        r.font.size = Pt(9.5)

        # Prominent Hall Banner (Bordered Box)
        banner_table = doc.add_table(rows=1, cols=1)
        banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        banner_table.autofit = False
        cell = banner_table.rows[0].cells[0]
        cell.width = Inches(7.3)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        # Background tint for the banner
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F2B48"/>')
        cell._tc.get_or_add_tcPr().append(shading)

        p_badge = cell.paragraphs[0]
        p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_badge.add_run("EXAMINATION VENUE & SEATING NOTICE")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(220, 235, 252)

        p_hall = cell.add_paragraph()
        p_hall.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_hall.paragraph_format.space_before = Pt(2)
        p_hall.paragraph_format.space_after = Pt(2)
        r = p_hall.add_run(alloc.venue_name.strip().upper())
        r.bold = True
        r.font.size = Pt(18)
        r.font.color.rgb = RGBColor(255, 255, 255)

        p_info = cell.add_paragraph()
        p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_info.add_run(f"COURSE: {config.course_title} ({config.course_code})   |   LEVEL: {config.level}")
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

        p_meta = cell.add_paragraph()
        p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p_meta.add_run(f"DATE: {config.exam_date}   |   ALLOCATED CANDIDATES: {len(alloc.students)}")
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(220, 235, 252)

        p_gap = doc.add_paragraph()
        p_gap.paragraph_format.space_before = Pt(4)
        p_gap.paragraph_format.space_after = Pt(4)

        # Instruction note for candidates
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(0)
        p_note.paragraph_format.space_after = Pt(4)
        r = p_note.add_run("NOTICE TO CANDIDATES: ")
        r.bold = True
        r.font.size = Pt(8.5)
        r2 = p_note.add_run("Only candidates listed below are authorized to enter and sit in this hall. Check your name and seat number.")
        r2.font.size = Pt(8.5)

        # 2-Column Roster Layout (to fit candidates efficiently on standard paper)
        # Each column has [S/N, REG NO, NAME]
        num_students = len(alloc.students)
        half = (num_students + 1) // 2

        roster_table = doc.add_table(rows=half + 1, cols=6)
        roster_table.style = 'Table Grid'
        roster_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        roster_table.autofit = False

        col_w = [Inches(0.4), Inches(1.3), Inches(1.95), Inches(0.4), Inches(1.3), Inches(1.95)]
        
        # Header Row
        hdr = roster_table.rows[0]
        headers = ["S/N", "REG NO", "STUDENT NAME", "S/N", "REG NO", "STUDENT NAME"]
        for c_idx, h_text in enumerate(headers):
            c = hdr.cells[c_idx]
            c.width = col_w[c_idx]
            set_cell_margins(c, top=60, bottom=60, left=50, right=50)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(8)

        # Student Rows
        for r_i in range(half):
            row = roster_table.rows[r_i + 1]
            left_s_idx = r_i
            right_s_idx = r_i + half

            # Left side
            if left_s_idx < num_students:
                s_left = alloc.students[left_s_idx]
                row.cells[0].paragraphs[0].add_run(str(left_s_idx + 1)).font.size = Pt(7.5)
                row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row.cells[1].paragraphs[0].add_run(s_left.reg_no).font.size = Pt(7.5)
                row.cells[2].paragraphs[0].add_run(s_left.name[:25]).font.size = Pt(7.5)

            # Right side
            if right_s_idx < num_students:
                s_right = alloc.students[right_s_idx]
                row.cells[3].paragraphs[0].add_run(str(right_s_idx + 1)).font.size = Pt(7.5)
                row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row.cells[4].paragraphs[0].add_run(s_right.reg_no).font.size = Pt(7.5)
                row.cells[5].paragraphs[0].add_run(s_right.name[:25]).font.size = Pt(7.5)

            for col_i in range(6):
                row.cells[col_i].width = col_w[col_i]
                set_cell_margins(row.cells[col_i], top=40, bottom=40, left=40, right=40)

    if output_target is None:
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    elif isinstance(output_target, str):
        doc.save(output_target)
        with open(output_target, 'rb') as f:
            return f.read()
    elif hasattr(output_target, 'write'):
        doc.save(output_target)
        if hasattr(output_target, 'getvalue'):
            return output_target.getvalue()
        return b''


def build_door_notices_pdf(
    config: ExamConfig,
    allocations: List[VenueAllocation],
    output_target: Union[str, io.BytesIO, None] = None
) -> bytes:
    """
    Generate official Hall Door Notices in PDF format.
    """
    buffer = io.BytesIO() if output_target is None or not isinstance(output_target, str) else output_target

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch
    )

    styles = getSampleStyleSheet()

    inst_style = ParagraphStyle(
        'NoticeInst',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        alignment=1
    )
    banner_title_style = ParagraphStyle(
        'BannerTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        alignment=1,
        textColor=colors.white
    )
    banner_sub_style = ParagraphStyle(
        'BannerSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#DCF0FF')
    )
    note_style = ParagraphStyle(
        'NoticeNote',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        alignment=0,
        textColor=colors.HexColor('#1E293B')
    )
    cell_center = ParagraphStyle(
        'CellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=8.5,
        alignment=1
    )
    cell_left = ParagraphStyle(
        'CellLeft',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=8.5,
        alignment=0
    )
    hdr_style = ParagraphStyle(
        'HdrStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=8.5,
        alignment=1
    )

    elements = []

    for hall_idx, alloc in enumerate(allocations):
        if hall_idx > 0:
            elements.append(PageBreak())

        # Header Titles
        elements.append(Paragraph(config.college.strip().upper(), inst_style))
        elements.append(Paragraph(config.department.strip().upper(), inst_style))
        elements.append(Paragraph(config.session_semester.strip(), inst_style))
        elements.append(Spacer(1, 4))

        # Big Venue Banner Box
        banner_content = [
            [Paragraph("EXAMINATION VENUE & SEATING NOTICE", banner_sub_style)],
            [Paragraph(alloc.venue_name.strip().upper(), banner_title_style)],
            [Paragraph(f"COURSE: {config.course_title} ({config.course_code})  |  LEVEL: {config.level}", banner_sub_style)],
            [Paragraph(f"DATE: {config.exam_date}  |  TOTAL CANDIDATES: {len(alloc.students)}", banner_sub_style)]
        ]
        banner_table = Table(banner_content, colWidths=[7.5 * inch])
        banner_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#0F2B48')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(banner_table)
        elements.append(Spacer(1, 4))

        # Notice to candidates
        elements.append(Paragraph(
            "<b>NOTICE:</b> Only candidates listed below are authorized to sit in this venue. Please verify your S/N and Registration Number before entering.",
            note_style
        ))
        elements.append(Spacer(1, 4))

        # 2-Column Roster Table
        num_students = len(alloc.students)
        half = (num_students + 1) // 2

        table_data = [[
            Paragraph("S/N", hdr_style),
            Paragraph("REG NO", hdr_style),
            Paragraph("NAME", hdr_style),
            Paragraph("S/N", hdr_style),
            Paragraph("REG NO", hdr_style),
            Paragraph("NAME", hdr_style)
        ]]

        for r_i in range(half):
            left_s_idx = r_i
            right_s_idx = r_i + half

            col_0 = str(left_s_idx + 1) if left_s_idx < num_students else ""
            col_1 = alloc.students[left_s_idx].reg_no if left_s_idx < num_students else ""
            col_2 = alloc.students[left_s_idx].name[:24] if left_s_idx < num_students else ""

            col_3 = str(right_s_idx + 1) if right_s_idx < num_students else ""
            col_4 = alloc.students[right_s_idx].reg_no if right_s_idx < num_students else ""
            col_5 = alloc.students[right_s_idx].name[:24] if right_s_idx < num_students else ""

            table_data.append([
                Paragraph(col_0, cell_center),
                Paragraph(col_1, cell_center),
                Paragraph(col_2, cell_left),
                Paragraph(col_3, cell_center),
                Paragraph(col_4, cell_center),
                Paragraph(col_5, cell_left)
            ])

        col_widths = [0.4 * inch, 1.35 * inch, 1.95 * inch, 0.4 * inch, 1.35 * inch, 1.95 * inch]
        roster_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        roster_table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(roster_table)

    doc.build(elements)

    if output_target is None:
        return buffer.getvalue()
    elif isinstance(output_target, str):
        with open(output_target, 'rb') as f:
            return f.read()
    elif hasattr(output_target, 'getvalue'):
        return output_target.getvalue()
    return b''
