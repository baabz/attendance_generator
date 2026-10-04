"""
Exam Materials Envelope Tag / Packet Docket Generator.
Produces official, printable envelope tags to attach to exam envelopes handed over to Central Exam Office for Chief Invigilators.
"""
import io
from typing import List, Union
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from .models import ExamConfig, VenueAllocation


def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    """Set inner padding for a table cell in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_envelope_tags_docx(
    config: ExamConfig,
    allocations: List[VenueAllocation],
    output_target: Union[str, io.BytesIO, None] = None
) -> bytes:
    """
    Generate official Envelope Tags in Word format.
    Each venue has a dedicated, professional docket sheet to paste on the packet envelope.
    """
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    for hall_idx, alloc in enumerate(allocations):
        if hall_idx > 0:
            doc.add_page_break()

        # Master Border Box (simulating docket envelope tag card)
        outer_table = doc.add_table(rows=1, cols=1)
        outer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        outer_table.autofit = False
        outer_cell = outer_table.rows[0].cells[0]
        outer_cell.width = Inches(7.3)
        set_cell_margins(outer_cell, top=140, bottom=140, left=180, right=180)

        # Institutional Header inside the tag
        p_inst = outer_cell.paragraphs[0]
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_inst.paragraph_format.space_before = Pt(0)
        p_inst.paragraph_format.space_after = Pt(2)
        r = p_inst.add_run(config.college.strip().upper())
        r.bold = True
        r.font.size = Pt(12)

        p_dept = outer_cell.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_dept.paragraph_format.space_before = Pt(0)
        p_dept.paragraph_format.space_after = Pt(2)
        r = p_dept.add_run(config.department.strip().upper())
        r.bold = True
        r.font.size = Pt(11)

        p_sem = outer_cell.add_paragraph()
        p_sem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sem.paragraph_format.space_before = Pt(0)
        p_sem.paragraph_format.space_after = Pt(6)
        r = p_sem.add_run(config.session_semester.strip())
        r.bold = True
        r.font.size = Pt(10)

        # Big Title Banner
        p_title = outer_cell.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(2)
        p_title.paragraph_format.space_after = Pt(8)
        r = p_title.add_run("EXAMINATION MATERIALS ENVELOPE TAG")
        r.bold = True
        r.font.size = Pt(14)
        r.font.underline = True

        # Candidate count and range calculation
        cand_count = len(alloc.students)
        first_reg = alloc.students[0].reg_no if alloc.students else "N/A"
        last_reg = alloc.students[-1].reg_no if alloc.students else "N/A"

        # Information Grid Table
        info_table = outer_cell.add_table(rows=7, cols=2)
        info_table.style = 'Table Grid'
        info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        info_table.autofit = False

        col_w_left = Inches(2.5)
        col_w_right = Inches(4.5)
        for r_item in info_table.rows:
            r_item.cells[0].width = col_w_left
            r_item.cells[1].width = col_w_right

        fields = [
            ("EXAM VENUE / HALL:", alloc.venue_name.strip().upper()),
            ("COURSE TITLE:", config.course_title.strip()),
            ("COURSE CODE & LEVEL:", f"{config.course_code.strip()}   |   {config.level.strip()}"),
            ("DATE & TIME:", f"{config.exam_date.strip()}   |   {config.time_in or '09:00 AM'} - {config.time_out or '12:00 PM'}"),
            ("TOTAL CANDIDATES:", f"{cand_count} CANDIDATES"),
            ("REG. NUMBER RANGE:", f"{first_reg}  TO  {last_reg}"),
            ("CHIEF INVIGILATOR:", config.chief_invigilator.strip() if config.chief_invigilator.strip() else "____________________________________")
        ]

        for idx, (label, val) in enumerate(fields):
            row_c = info_table.rows[idx]
            set_cell_margins(row_c.cells[0], top=60, bottom=60, left=80, right=80)
            set_cell_margins(row_c.cells[1], top=60, bottom=60, left=80, right=80)

            p_lbl = row_c.cells[0].paragraphs[0]
            p_lbl.paragraph_format.space_before = Pt(0)
            p_lbl.paragraph_format.space_after = Pt(0)
            r_lbl = p_lbl.add_run(label)
            r_lbl.bold = True
            r_lbl.font.size = Pt(9.5)

            p_val = row_c.cells[1].paragraphs[0]
            p_val.paragraph_format.space_before = Pt(0)
            p_val.paragraph_format.space_after = Pt(0)
            r_val = p_val.add_run(val)
            r_val.bold = True if idx in [0, 4] else False
            r_val.font.size = Pt(11 if idx == 0 else 9.5)

        # Enclosed Materials Checklist
        p_chk_title = outer_cell.add_paragraph()
        p_chk_title.paragraph_format.space_before = Pt(8)
        p_chk_title.paragraph_format.space_after = Pt(2)
        r = p_chk_title.add_run("ENCLOSED EXAMINATION MATERIALS (VERIFY BEFORE DEPARTURE):")
        r.bold = True
        r.font.size = Pt(9.5)

        chk_table = outer_cell.add_table(rows=5, cols=3)
        chk_table.style = 'Table Grid'
        chk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        chk_table.autofit = False

        chk_items = [
            ("ITEM DESCRIPTION", "QUANTITY ENCLOSED", "INVIGILATOR CHECK"),
            (f"1. Question Papers ({config.course_code})", f"{cand_count} Copies", "[   ] Verified"),
            (f"2. Examination Answer Booklets", f"{cand_count} Booklets", "[   ] Verified"),
            (f"3. Official Attendance Register (Hall Copy)", "1 Copy (Enclosed)", "[   ] Verified"),
            (f"4. Entrance Door Notice Slip", "1 Copy (Enclosed)", "[   ] Verified")
        ]

        chk_widths = [Inches(3.8), Inches(1.8), Inches(1.4)]
        for r_idx, row_data in enumerate(chk_items):
            c_row = chk_table.rows[r_idx]
            for col_i in range(3):
                c_cell = c_row.cells[col_i]
                c_cell.width = chk_widths[col_i]
                set_cell_margins(c_cell, top=45, bottom=45, left=60, right=60)
                p = c_cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                if col_i in [1, 2]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_txt = p.add_run(row_data[col_i])
                r_txt.font.size = Pt(8.5)
                if r_idx == 0:
                    r_txt.bold = True

        # Handover Signatures & Chain of Custody
        p_sign_title = outer_cell.add_paragraph()
        p_sign_title.paragraph_format.space_before = Pt(8)
        p_sign_title.paragraph_format.space_after = Pt(2)
        r = p_sign_title.add_run("CHAIN OF CUSTODY & HANDOVER LOG:")
        r.bold = True
        r.font.size = Pt(9.5)

        sign_table = outer_cell.add_table(rows=3, cols=3)
        sign_table.style = 'Table Grid'
        sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sign_table.autofit = False

        sign_rows = [
            ("STAGE / OPERATION", "OFFICER NAME & DESIGNATION", "SIGNATURE & TIME"),
            ("1. DISPATCHED BY (Dept Exam Officer):", "........................................................", "Sign: ................... Time: ........"),
            ("2. PICKED UP BY (Chief Invigilator):", "........................................................", "Sign: ................... Time: ........")
        ]

        s_widths = [Inches(2.5), Inches(2.6), Inches(1.9)]
        for r_idx, s_data in enumerate(sign_rows):
            s_row = sign_table.rows[r_idx]
            for c_i in range(3):
                sc = s_row.cells[c_i]
                sc.width = s_widths[c_i]
                set_cell_margins(sc, top=50, bottom=50, left=60, right=60)
                p = sc.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                r_t = p.add_run(s_data[c_i])
                r_t.font.size = Pt(8)
                if r_idx == 0:
                    r_t.bold = True

    if output_target is None:
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    elif isinstance(output_target, str):
        doc.save(output_target)
        with open(output_target, 'rb') as f:
            return f.read()
    elif hasattr(output_target, 'write'):
        doc.save(output_target)
        if hasattr(output_target, 'getvalue'):
            return output_target.getvalue()
        return b''


def build_envelope_tags_pdf(
    config: ExamConfig,
    allocations: List[VenueAllocation],
    output_target: Union[str, io.BytesIO, None] = None
) -> bytes:
    """
    Generate official Envelope Tags in PDF format.
    """
    buffer = io.BytesIO() if output_target is None or not isinstance(output_target, str) else output_target

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.4 * inch,
        bottomMargin=0.4 * inch
    )

    styles = getSampleStyleSheet()

    inst_style = ParagraphStyle(
        'TagInst',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=13,
        alignment=1
    )
    banner_title_style = ParagraphStyle(
        'TagTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor('#0F2B48')
    )
    label_style = ParagraphStyle(
        'TagLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=0
    )
    val_style = ParagraphStyle(
        'TagVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=0
    )
    val_bold_style = ParagraphStyle(
        'TagValBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        alignment=0,
        textColor=colors.HexColor('#0F2B48')
    )
    section_hdr_style = ParagraphStyle(
        'TagSecHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        alignment=0,
        textColor=colors.HexColor('#1E293B')
    )
    table_cell = ParagraphStyle(
        'TagTblCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=0
    )
    table_cell_center = ParagraphStyle(
        'TagTblCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1
    )
    table_cell_hdr = ParagraphStyle(
        'TagTblHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        alignment=1
    )

    elements = []

    for hall_idx, alloc in enumerate(allocations):
        if hall_idx > 0:
            elements.append(PageBreak())

        cand_count = len(alloc.students)
        first_reg = alloc.students[0].reg_no if alloc.students else "N/A"
        last_reg = alloc.students[-1].reg_no if alloc.students else "N/A"

        tag_elements = []

        # Header Titles
        tag_elements.append(Paragraph(config.college.strip().upper(), inst_style))
        tag_elements.append(Paragraph(config.department.strip().upper(), inst_style))
        tag_elements.append(Paragraph(config.session_semester.strip(), inst_style))
        tag_elements.append(Spacer(1, 4))
        tag_elements.append(Paragraph("<u>EXAMINATION MATERIALS ENVELOPE TAG</u>", banner_title_style))
        tag_elements.append(Spacer(1, 6))

        # Metadata Table
        fields_data = [
            [Paragraph("EXAM VENUE / HALL:", label_style), Paragraph(alloc.venue_name.strip().upper(), val_bold_style)],
            [Paragraph("COURSE TITLE:", label_style), Paragraph(config.course_title.strip(), val_style)],
            [Paragraph("COURSE CODE & LEVEL:", label_style), Paragraph(f"{config.course_code.strip()}  |  {config.level.strip()}", val_style)],
            [Paragraph("DATE & TIME:", label_style), Paragraph(f"{config.exam_date.strip()}  |  {config.time_in or '09:00 AM'} - {config.time_out or '12:00 PM'}", val_style)],
            [Paragraph("TOTAL CANDIDATES:", label_style), Paragraph(f"<b>{cand_count} CANDIDATES</b>", val_style)],
            [Paragraph("REG. NUMBER RANGE:", label_style), Paragraph(f"<b>{first_reg}  TO  {last_reg}</b>", val_style)],
            [Paragraph("CHIEF INVIGILATOR:", label_style), Paragraph(config.chief_invigilator.strip() if config.chief_invigilator.strip() else "____________________________________", val_style)]
        ]

        t_meta = Table(fields_data, colWidths=[2.3 * inch, 4.7 * inch])
        t_meta.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#64748B')),
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F8FAFC')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        tag_elements.append(t_meta)
        tag_elements.append(Spacer(1, 8))

        # Enclosed Checklist Table
        tag_elements.append(Paragraph("<b>ENCLOSED MATERIALS CHECKLIST (VERIFY BEFORE PICKUP):</b>", section_hdr_style))
        tag_elements.append(Spacer(1, 3))

        chk_data = [
            [Paragraph("ITEM DESCRIPTION", table_cell_hdr), Paragraph("QUANTITY ENCLOSED", table_cell_hdr), Paragraph("CHECK", table_cell_hdr)],
            [Paragraph(f"1. Question Papers ({config.course_code})", table_cell), Paragraph(f"{cand_count} Copies", table_cell_center), Paragraph("[   ] Verified", table_cell_center)],
            [Paragraph("2. Examination Answer Booklets", table_cell), Paragraph(f"{cand_count} Booklets", table_cell_center), Paragraph("[   ] Verified", table_cell_center)],
            [Paragraph("3. Official Attendance Register (Hall Copy)", table_cell), Paragraph("1 Copy (Enclosed)", table_cell_center), Paragraph("[   ] Verified", table_cell_center)],
            [Paragraph("4. Venue Entrance Door Notice", table_cell), Paragraph("1 Copy (Enclosed)", table_cell_center), Paragraph("[   ] Verified", table_cell_center)],
        ]
        t_chk = Table(chk_data, colWidths=[3.8 * inch, 1.8 * inch, 1.4 * inch])
        t_chk.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#64748B')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        tag_elements.append(t_chk)
        tag_elements.append(Spacer(1, 8))

        # Handover Table
        tag_elements.append(Paragraph("<b>CHAIN OF CUSTODY & HANDOVER LOG:</b>", section_hdr_style))
        tag_elements.append(Spacer(1, 3))

        sign_data = [
            [Paragraph("STAGE / OPERATION", table_cell_hdr), Paragraph("OFFICER NAME & DESIGNATION", table_cell_hdr), Paragraph("SIGNATURE & TIME", table_cell_hdr)],
            [
                Paragraph("1. DISPATCHED BY (Dept Exam Officer):", table_cell),
                Paragraph("........................................................", table_cell),
                Paragraph("Sign: ............. Time: .....", table_cell)
            ],
            [
                Paragraph("2. RECEIVED BY (Chief Invigilator):", table_cell),
                Paragraph("........................................................", table_cell),
                Paragraph("Sign: ............. Time: .....", table_cell)
            ]
        ]
        t_sign = Table(sign_data, colWidths=[2.5 * inch, 2.7 * inch, 1.8 * inch])
        t_sign.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#64748B')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        tag_elements.append(t_sign)

        # Wrap in an outer bordered card
        outer_card = Table([[tag_elements]], colWidths=[7.4 * inch])
        outer_card.setStyle(TableStyle([
            ('BOX', (0, 0), (-1, -1), 1.5, colors.HexColor('#0F2B48')),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
            ('RIGHTPADDING', (0, 0), (-1, -1), 12),
        ]))

        elements.append(outer_card)

    doc.build(elements)

    if output_target is None:
        return buffer.getvalue()
    elif isinstance(output_target, str):
        with open(output_target, 'rb') as f:
            return f.read()
    elif hasattr(output_target, 'getvalue'):
        return output_target.getvalue()
    return b''
